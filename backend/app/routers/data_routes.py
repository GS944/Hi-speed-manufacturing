"""Data sheets: uploads, storage, tables, rows, mapping editor, exports, data quality."""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
import shutil
import threading
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from .. import config
from ..auth import require_admin
from ..db import AuditLog, DataFile, DataTable, Record, User, get_db
from ..ingest import pipeline
from ..intelligence import column_model, profiler
from ..services import storage, workspace
from ..services.keys import to_date, to_num

router = APIRouter(prefix="/api", tags=["data"])
log = logging.getLogger("ordertrack.data")
ALLOWED = {".xlsx", ".xls", ".xlsm", ".csv"}


# ----------------------------------------------------------------------------- storage
def _database_bytes(db: Session) -> int:
    if config.IS_SQLITE:
        return sum(p.stat().st_size for p in config.DATA_DIR.glob("ordertrack.db*") if p.is_file())
    try:
        return int(db.execute(text("SELECT pg_database_size(current_database())")).scalar() or 0)
    except Exception:  # noqa: BLE001
        return 0


def storage_usage(db: Session) -> dict:
    files = int(db.scalar(select(func.coalesce(func.sum(DataFile.size_bytes), 0))) or 0)
    dbsize = _database_bytes(db)
    quota = int(config.STORAGE_QUOTA_GB * 1024 ** 3)
    used = files + dbsize
    local = config.STORAGE_BACKEND == "local"
    disk_free = shutil.disk_usage(config.DATA_DIR).free if local else None
    available = max(0, quota - used) if not local else max(0, min(quota - used, disk_free))
    return {"used_bytes": used, "files_bytes": files, "database_bytes": dbsize,
            "models_bytes": sum(p.stat().st_size for p in config.MODELS_DIR.glob("*") if p.is_file()),
            "quota_bytes": quota, "quota_gb": config.STORAGE_QUOTA_GB, "disk_free_bytes": disk_free,
            "available_bytes": available, "max_upload_mb": config.MAX_UPLOAD_MB,
            "backend": "S3-compatible bucket" if not local else "Local disk",
            "database": "SQLite" if config.IS_SQLITE else "PostgreSQL",
            "used_pct": round(100 * used / quota, 2) if quota else 0}


@router.get("/storage")
def get_storage(_: User = Depends(require_admin), db: Session = Depends(get_db)):
    return storage_usage(db)


# ----------------------------------------------------------------------------- files
def _file_json(f: DataFile, tables: list[DataTable] | None = None) -> dict:
    out = {"id": f.id, "name": f.original_name, "size_bytes": f.size_bytes, "status": f.status, "error": f.error,
           "uploaded_by": f.uploaded_by, "uploaded_at": f.uploaded_at.isoformat(), "notes": f.notes,
           "processed_at": f.processed_at.isoformat() if f.processed_at else None, "sha256": f.sha256}
    if tables is not None:
        out["tables"] = [_table_brief(t) for t in tables]
    return out


def _table_brief(t: DataTable) -> dict:
    q = workspace.get().quality.get(t.id) if t.kind == "table" and t.active else None
    return {"id": t.id, "file_id": t.file_id, "title": t.title, "sheet": t.sheet_name, "range": t.range_ref, "kind": t.kind,
            "role": t.role, "role_label": (t.profile or {}).get("type_label") or ("Print template" if t.kind == "form" else t.role),
            "confidence": t.role_confidence, "rows": t.n_rows, "columns": len(t.columns or []), "active": t.active,
            "quality_score": q["score"] if q else None, "issue_count": sum(i["count"] for i in q["issues"] if i["severity"] != "info") if q else 0}


@router.get("/files")
def list_files(_: User = Depends(require_admin), db: Session = Depends(get_db)):
    files = db.scalars(select(DataFile).order_by(DataFile.uploaded_at.desc())).all()
    tables = db.scalars(select(DataTable).order_by(DataTable.sheet_name, DataTable.block_index)).all()
    by_file: dict[int, list] = {}
    for t in tables:
        by_file.setdefault(t.file_id, []).append(t)
    return [_file_json(f, by_file.get(f.id, [])) for f in files]


def _sheet_signature(path: Path) -> tuple[str, ...]:
    try:
        if path.suffix.lower() == ".xls":
            import xlrd
            return tuple(sorted(xlrd.open_workbook(str(path), on_demand=True).sheet_names()))
        if path.suffix.lower() in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True)
            names = tuple(sorted(wb.sheetnames))
            wb.close()
            return names
    except Exception:  # noqa: BLE001 - a broken file is reported later by the pipeline
        pass
    return ()


@router.post("/files")
async def upload_file(file: UploadFile = File(...), replace_file_id: int | None = Form(None), notes: str | None = Form(None),
                      keep_both: bool = Form(False),
                      user: User = Depends(require_admin), db: Session = Depends(get_db)):
    name = Path(file.filename or "upload").name
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED:
        raise HTTPException(400, f"Unsupported file type {ext or '(none)'}. Allowed: {', '.join(sorted(ALLOWED))}")
    usage = storage_usage(db)
    stored = f"{int(time.time() * 1000)}_{re.sub(r'[^A-Za-z0-9._-]+', '_', name)}"
    dest = config.TMP_DIR / stored            # staged locally, validated, then moved to permanent storage
    h, size, limit = hashlib.sha256(), 0, config.MAX_UPLOAD_MB * 1024 * 1024
    try:
        with dest.open("wb") as out:
            while chunk := await file.read(1024 * 1024):
                size += len(chunk)
                if size > limit:
                    raise HTTPException(413, f"File exceeds the {config.MAX_UPLOAD_MB} MB upload limit.")
                if size > usage["available_bytes"]:
                    raise HTTPException(507, "Storage quota exceeded. Delete old files or raise STORAGE_QUOTA_GB.")
                h.update(chunk)
                out.write(chunk)
    except HTTPException:
        dest.unlink(missing_ok=True)
        raise
    digest = h.hexdigest()
    dup = db.scalar(select(DataFile).where(DataFile.sha256 == digest, DataFile.status != "archived"))
    if dup and not replace_file_id:
        dest.unlink(missing_ok=True)
        raise HTTPException(409, f"This exact file is already loaded as “{dup.original_name}”.")
    sig = "|".join(_sheet_signature(dest)) or None
    if not replace_file_id and not keep_both:
        # same workbook re-saved (different bytes) would double-count every order - ask first
        for other in db.scalars(select(DataFile).where(DataFile.status.in_(["ready", "processing"]))):
            if other.sheet_sig is None:                  # uploaded before signatures were stored: backfill once
                try:
                    with storage.get().local_copy(other.stored_name) as p:
                        other.sheet_sig = "|".join(_sheet_signature(p))
                except Exception:  # noqa: BLE001
                    other.sheet_sig = ""
            same_name = other.original_name.lower() == name.lower()
            same_shape = bool(sig) and sig == other.sheet_sig
            if same_name or same_shape:
                dest.unlink(missing_ok=True)
                raise HTTPException(409, {"code": "possible_new_version", "file_id": other.id, "file_name": other.original_name,
                                          "message": f"This looks like a new version of “{other.original_name}”. Replace it, or keep both?"})
    try:
        storage.get().put(dest, stored)
    except Exception:
        dest.unlink(missing_ok=True)
        log.exception("could not store upload %s", name)
        raise HTTPException(503, "The file could not be saved to storage. Please try again in a minute.")
    f = DataFile(original_name=name, stored_name=stored, size_bytes=size, sha256=digest, sheet_sig=sig,
                 uploaded_by=user.username, notes=notes)
    db.add(f)
    if replace_file_id:
        old = db.get(DataFile, replace_file_id)
        if old:
            old.status = "archived"
            for t in db.scalars(select(DataTable).where(DataTable.file_id == old.id)):
                t.active = False
    db.add(AuditLog(username=user.username, action="upload", entity="file", detail={"name": name, "size": size, "replaces": replace_file_id}))
    db.commit()
    pipeline.process_file_async(f.id)
    return _file_json(f)


@router.get("/files/{file_id}")
def get_file(file_id: int, _: User = Depends(require_admin), db: Session = Depends(get_db)):
    f = db.get(DataFile, file_id) or _404("File")
    tables = db.scalars(select(DataTable).where(DataTable.file_id == f.id).order_by(DataTable.sheet_name, DataTable.block_index)).all()
    return _file_json(f, tables)


@router.get("/files/{file_id}/download")
def download_file(file_id: int, _: User = Depends(require_admin), db: Session = Depends(get_db)):
    f = db.get(DataFile, file_id) or _404("File")
    return StreamingResponse(storage.get().iter_bytes(f.stored_name), media_type="application/octet-stream",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(f.original_name)}",
                                      "Content-Length": str(f.size_bytes)})


@router.post("/files/{file_id}/reprocess")
def reprocess(file_id: int, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    f = db.get(DataFile, file_id) or _404("File")
    if f.status == "archived":
        raise HTTPException(400, "Restore the file before reprocessing it.")
    f.status = "processing"
    db.add(AuditLog(username=user.username, action="reprocess", entity="file", entity_id=str(file_id)))
    db.commit()
    pipeline.process_file_async(f.id)
    return {"ok": True}


@router.post("/files/{file_id}/archive")
def archive(file_id: int, restore: bool = False, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    f = db.get(DataFile, file_id) or _404("File")
    f.status = "ready" if restore else "archived"
    for t in db.scalars(select(DataTable).where(DataTable.file_id == f.id)):
        t.active = restore
    db.add(AuditLog(username=user.username, action="restore" if restore else "archive", entity="file", entity_id=str(file_id)))
    db.commit()
    workspace.invalidate()
    return _file_json(f)


@router.delete("/files/{file_id}")
def delete_file(file_id: int, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    f = db.get(DataFile, file_id) or _404("File")
    try:
        storage.get().delete(f.stored_name)
    except Exception:  # noqa: BLE001 - the database row is still removed; log the orphan object
        log.exception("could not delete stored object %s", f.stored_name)
    db.delete(f)
    db.add(AuditLog(username=user.username, action="delete", entity="file", detail={"name": f.original_name}))
    db.commit()
    workspace.invalidate()
    return {"ok": True}


# ----------------------------------------------------------------------------- tables
@router.get("/tables")
def list_tables(role: str | None = None, _: User = Depends(require_admin), db: Session = Depends(get_db)):
    stmt = select(DataTable).join(DataFile).where(DataTable.active.is_(True), DataFile.status == "ready")
    if role:
        stmt = stmt.where(DataTable.role == role)
    return [_table_brief(t) | {"file": f.original_name} for t, f in
            db.execute(stmt.add_columns(DataFile).order_by(DataFile.uploaded_at.desc(), DataTable.sheet_name)).all()]


@router.get("/tables/{table_id}")
def get_table(table_id: int, _: User = Depends(require_admin), db: Session = Depends(get_db)):
    t = db.get(DataTable, table_id) or _404("Table")
    f = db.get(DataFile, t.file_id)
    return _table_brief(t) | {"file": f.original_name if f else "", "columns": t.columns, "profile": t.profile,
                              "header_row": t.header_row, "form": t.form, "quality": workspace.get().quality.get(t.id)}


def _cell_text(v) -> str:
    return "" if v is None else str(v)


@router.get("/tables/{table_id}/rows")
def table_rows(table_id: int, page: int = Query(1, ge=1), size: int = Query(50, ge=1, le=500), q: str | None = None,
               sort: str | None = None, dir: str = "asc", rows: str | None = None,
               _: User = Depends(require_admin), db: Session = Depends(get_db)):
    t = db.get(DataTable, table_id) or _404("Table")
    recs = db.execute(select(Record.id, Record.row_no, Record.data).where(Record.table_id == t.id)).all()
    items = [{"id": r.id, "row_no": r.row_no, "data": r.data} for r in recs]
    if rows:
        wanted = {int(x) for x in rows.split(",") if x.strip().isdigit()}
        items = [i for i in items if i["row_no"] in wanted]
    if q:
        ql = q.lower()
        items = [i for i in items if any(ql in _cell_text(v).lower() for v in i["data"].values())]
    if sort:
        dtype = next((c["dtype"] for c in t.columns if c["name"] == sort), "text")

        def key(i):
            v = i["data"].get(sort)
            if v is None:
                return (1, 0, "")
            if dtype in ("integer", "number"):
                n = to_num(v)
                return (0, n if n is not None else 0, "")
            return (0, 0, str(v).lower())
        items.sort(key=key, reverse=dir == "desc")
        if dir == "desc":   # keep blanks last
            items.sort(key=lambda i: i["data"].get(sort) is None)
    else:
        items.sort(key=lambda i: (i["row_no"] == 0, i["row_no"], i["id"]))
    total = len(items)
    start = (page - 1) * size
    return {"total": total, "page": page, "size": size, "rows": items[start:start + size]}


class RowIn(BaseModel):
    data: dict


def _coerce(t: DataTable, data: dict) -> dict:
    dtypes = {c["name"]: c["dtype"] for c in t.columns}
    out = {}
    for k, v in data.items():
        if k not in dtypes:
            raise HTTPException(400, f"Unknown column '{k}'.")
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        dt = dtypes[k]
        if dt in ("integer", "number"):
            n = to_num(v)
            if n is None:
                raise HTTPException(400, f"'{k}' must be a number.")
            out[k] = int(n) if n.is_integer() else n
        elif dt == "date":
            d = to_date(v)
            if d is None:
                raise HTTPException(400, f"'{k}' must be a date (YYYY-MM-DD).")
            out[k] = d.isoformat()
        else:
            out[k] = str(v).strip() if isinstance(v, str) else v
    return out


@router.post("/tables/{table_id}/rows")
def add_row(table_id: int, body: RowIn, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    t = db.get(DataTable, table_id) or _404("Table")
    rec = Record(table_id=t.id, row_no=0, data=_coerce(t, body.data))
    db.add(rec)
    db.flush()
    t.n_rows += 1
    pipeline.reindex_record(db, t, rec)
    db.add(AuditLog(username=user.username, action="add_row", entity="table", entity_id=str(t.id), detail={"record": rec.id, "data": rec.data}))
    db.commit()
    workspace.invalidate()
    return {"id": rec.id, "row_no": rec.row_no, "data": rec.data}


@router.put("/tables/{table_id}/rows/{record_id}")
def update_row(table_id: int, record_id: int, body: RowIn, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    t = db.get(DataTable, table_id) or _404("Table")
    rec = db.get(Record, record_id)
    if not rec or rec.table_id != t.id:
        _404("Row")
    before = dict(rec.data)
    merged = dict(rec.data)
    for k, v in body.data.items():
        merged.pop(k, None)
    merged.update(_coerce(t, body.data))
    rec.data = merged
    pipeline.reindex_record(db, t, rec)
    changes = {k: {"from": before.get(k), "to": merged.get(k)} for k in body.data if before.get(k) != merged.get(k)}
    db.add(AuditLog(username=user.username, action="edit_row", entity="table", entity_id=str(t.id),
                    detail={"record": rec.id, "row_no": rec.row_no, "changes": changes}))
    db.commit()
    workspace.invalidate()
    return {"id": rec.id, "row_no": rec.row_no, "data": rec.data}


@router.delete("/tables/{table_id}/rows/{record_id}")
def delete_row(table_id: int, record_id: int, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    t = db.get(DataTable, table_id) or _404("Table")
    rec = db.get(Record, record_id)
    if not rec or rec.table_id != t.id:
        _404("Row")
    db.add(AuditLog(username=user.username, action="delete_row", entity="table", entity_id=str(t.id), detail={"record": rec.id, "data": rec.data}))
    db.delete(rec)
    t.n_rows = max(0, t.n_rows - 1)
    db.commit()
    workspace.invalidate()
    return {"ok": True}


@router.get("/tables/{table_id}/export")
def export_table(table_id: int, format: str = "xlsx", _: User = Depends(require_admin), db: Session = Depends(get_db)):
    t = db.get(DataTable, table_id) or _404("Table")
    cols = [c["name"] for c in t.columns]
    recs = db.scalars(select(Record).where(Record.table_id == t.id).order_by(Record.row_no == 0, Record.row_no, Record.id)).all()
    fname = re.sub(r"[^A-Za-z0-9 _-]+", "", t.sheet_name).strip() or "export"
    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(cols)
        for r in recs:
            w.writerow([r.data.get(c, "") for c in cols])
        return Response(buf.getvalue().encode("utf-8-sig"), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{fname}.csv"'})
    wb = Workbook()
    ws = wb.active
    ws.title = fname[:31]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E2E8F0")
    dtypes = {c["name"]: c["dtype"] for c in t.columns}
    for r in recs:
        row = []
        for c in cols:
            v = r.data.get(c)
            if dtypes.get(c) == "date" and isinstance(v, str):
                d = to_date(v)
                v = datetime(d.year, d.month, d.day) if d else v
            row.append(v)
        ws.append(row)
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(45, max(10, len(c) + 4))
        if dtypes.get(c) == "date":
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for x in cell:
                    x.number_format = "dd-mm-yyyy"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'})


# ----------------------------------------------------------------------------- mapping editor
class ProfileIn(BaseModel):
    table_type: str
    fields: dict[str, str | None]
    title: str | None = None
    overrides: dict | None = None      # advanced: stages / completion / urgent_values JSON


@router.get("/ontology")
def ontology(_: User = Depends(require_admin)):
    o = column_model.load_ontology()
    return {"roles": [{"key": k, "label": v["label"], "kind": v["kind"]} for k, v in o["roles"].items()],
            "table_types": [{"key": k, "label": v["label"]} for k, v in o["table_types"].items()] +
                           [{"key": "unknown", "label": "Unclassified"}]}


@router.put("/tables/{table_id}/profile")
def update_profile(table_id: int, body: ProfileIn, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    t = db.get(DataTable, table_id) or _404("Table")
    if t.kind != "table":
        raise HTTPException(400, "Print templates have no column mapping.")
    onto = column_model.load_ontology()
    names = {c["name"] for c in t.columns}
    fields = {r: c for r, c in body.fields.items() if c}
    bad = [c for c in fields.values() if c not in names]
    if bad:
        raise HTTPException(400, f"Unknown column(s): {', '.join(bad)}")
    if body.table_type not in onto["table_types"] and body.table_type != "unknown":
        raise HTTPException(400, "Unknown table type.")
    if len(set(fields.values())) != len(fields):
        raise HTTPException(400, "A column can only be mapped to one role.")
    old = t.profile or {}
    all_recs = db.scalars(select(Record).where(Record.table_id == t.id)).all()
    values = {c: [r.data.get(c) for r in all_recs] for c in names}
    conf = {r: (1.0 if old.get("fields", {}).get(r) != c else old.get("field_confidence", {}).get(r, 1.0)) for r, c in fields.items()}
    col_role = {c: r for r, c in fields.items()}
    per_col = [{"column": c["name"], "role": col_role.get(c["name"]), "confidence": conf.get(col_role.get(c["name"]), 0) if col_role.get(c["name"]) else 0,
                "alternatives": next((pc["alternatives"] for pc in old.get("columns", []) if pc["column"] == c["name"]), [])}
               for c in t.columns]
    prof = profiler.build_profile(body.table_type, 1.0, fields, conf, per_col, values, t.sheet_name, onto, old.get("type_scores", {}))
    prof["edited_by"], prof["edited_at"] = user.username, datetime.now().isoformat(timespec="seconds")
    if body.overrides:
        for k in ("stages", "completion", "urgent_values"):
            if k in body.overrides:
                prof[k] = body.overrides[k]
    t.profile = prof
    t.role = body.table_type
    t.role_confidence = 1.0
    t.title = body.title or profiler.title_for(body.table_type, t.sheet_name, onto)
    pipeline.reindex_table(db, t)
    # learn from the admin: every mapped column becomes a labelled training example
    _parse = lambda v: datetime.fromisoformat(v) if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) else v
    examples = [(c, [_parse(v) for v in values[c][:400] if v is not None][:80], r) for r, c in fields.items()]
    changed = {r: c for r, c in fields.items() if old.get("fields", {}).get(r) != c}
    db.add(AuditLog(username=user.username, action="edit_mapping", entity="table", entity_id=str(t.id),
                    detail={"table_type": body.table_type, "changed": changed}))
    db.commit()
    column_model.add_feedback(examples)
    threading.Thread(target=column_model.train, daemon=True).start()
    workspace.invalidate()
    return get_table(table_id, user, db)


@router.post("/tables/{table_id}/toggle")
def toggle_table(table_id: int, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    t = db.get(DataTable, table_id) or _404("Table")
    t.active = not t.active
    db.add(AuditLog(username=user.username, action="toggle_table", entity="table", entity_id=str(t.id), detail={"active": t.active}))
    db.commit()
    workspace.invalidate()
    return _table_brief(t)


@router.get("/quality")
def quality(_: User = Depends(require_admin), db: Session = Depends(get_db)):
    ws = workspace.get()
    out = []
    for tid, q in ws.quality.items():
        t = ws.tables.get(tid)
        if t and t["role"] in ("orders", "job_work_ledger", "customer_master", "vendor_master", "price_list"):
            out.append(q | {"title": t["title"], "file": t["file"], "role": t["role"]})
    return sorted(out, key=lambda x: x["score"])


def _404(what: str):
    raise HTTPException(404, f"{what} not found.")
