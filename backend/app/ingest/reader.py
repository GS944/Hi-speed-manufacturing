"""Workbook reader + table segmentation.

Nothing here knows about orders or annexures. A sheet is treated as a 2-D grid; the grid is
segmented into independent blocks (runs of non-empty columns separated by empty columns), and in
each block the header row is located by a scoring function. Blocks that behave like printable
forms (repeated headers, label/value layouts) are kept as *forms* instead of tables.
"""
from __future__ import annotations

import csv
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter

Cell = Any


@dataclass
class RawSheet:
    name: str
    grid: list[list[Cell]]
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)  # r1,c1,r2,c2 (0-based, inclusive)


@dataclass
class RawTable:
    sheet: str
    block_index: int
    kind: str                      # table | form
    first_col: int
    last_col: int
    header_row: int                # 0-based within sheet
    columns: list[str]
    letters: list[str]
    rows: list[list[Cell]]         # data rows (python values)
    row_numbers: list[int]         # 1-based sheet row numbers
    range_ref: str
    form_cells: list[tuple[int, int, Cell]] | None = None
    merged: list[tuple[int, int, int, int]] | None = None


# ----------------------------------------------------------------------------- cell normalisation
_WS = re.compile(r"\s+")


def clean(v: Cell) -> Cell:
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return None
        return int(v) if v.is_integer() and abs(v) < 1e15 else v
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, str):
        s = _WS.sub(" ", v.replace(" ", " ")).strip()
        if not s or s.startswith("#") and s.rstrip("!").upper() in {"#N/A", "#REF", "#VALUE", "#DIV/0", "#NAME?", "#NUM"}:
            return None
        return s
    return v


def is_empty(v: Cell) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


# ----------------------------------------------------------------------------- readers
def read_workbook(path: Path) -> list[RawSheet]:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        return _read_openpyxl(path)
    if ext == ".xls":
        return _read_xlrd(path)
    if ext in (".csv", ".tsv", ".txt"):
        return _read_csv(path)
    raise ValueError(f"Unsupported file type '{ext}'. Upload .xlsx, .xls, .xlsm or .csv")


def _trim(grid: list[list[Cell]]) -> list[list[Cell]]:
    last_r, last_c = -1, -1
    for r, row in enumerate(grid):
        for c, v in enumerate(row):
            if not is_empty(v):
                last_r = r
                last_c = max(last_c, c)
    return [list(row[: last_c + 1]) + [None] * max(0, last_c + 1 - len(row)) for row in grid[: last_r + 1]]


def _read_openpyxl(path: Path) -> list[RawSheet]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    out = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" and ws.max_row <= 1:
            continue
        grid = [[clean(v) for v in row] for row in ws.iter_rows(values_only=True)]
        merged = [(m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1) for m in ws.merged_cells.ranges]
        out.append(RawSheet(ws.title, _trim(grid), merged))
    wb.close()
    return out


def _read_xlrd(path: Path) -> list[RawSheet]:
    import xlrd
    try:
        wb = xlrd.open_workbook(str(path), formatting_info=True)
    except NotImplementedError:
        wb = xlrd.open_workbook(str(path))
    out = []
    for sh in wb.sheets():
        grid = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate_as_datetime(cell.value, wb.datemode))
                    except Exception:
                        row.append(clean(cell.value))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    row.append(None)
                else:
                    row.append(clean(cell.value))
            grid.append(row)
        merged = [(r1, c1, r2 - 1, c2 - 1) for (r1, r2, c1, c2) in getattr(sh, "merged_cells", [])]
        out.append(RawSheet(sh.name, _trim(grid), merged))
    return out


def _read_csv(path: Path) -> list[RawSheet]:
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|") if text else csv.excel
    rows = [[clean(_num(v)) for v in r] for r in csv.reader(text.splitlines(), dialect)]
    return [RawSheet(path.stem, _trim(rows))]


def _num(v: str):
    s = v.strip()
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d*\.\d+", s):
        return float(s)
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return v


# ----------------------------------------------------------------------------- segmentation
def _col_blocks(grid: list[list[Cell]], merged=()) -> list[tuple[int, int]]:
    if not grid:
        return []
    ncols = max(len(r) for r in grid)
    occupied = [False] * ncols
    for row in grid:
        for c, v in enumerate(row):
            if not is_empty(v):
                occupied[c] = True
    # a merged cell with content bridges the columns it spans (e.g. a description merged over C:E)
    for r1, c1, _r2, c2 in merged:
        if r1 < len(grid) and c1 < len(grid[r1]) and not is_empty(grid[r1][c1]):
            for c in range(c1, min(c2, ncols - 1) + 1):
                occupied[c] = True
    blocks, start = [], None
    for c, occ in enumerate(occupied + [False]):
        if occ and start is None:
            start = c
        elif not occ and start is not None:
            blocks.append((start, c - 1))
            start = None
    return blocks


_LABELISH = re.compile(r"^[A-Za-z][A-Za-z .&/()'-]{1,40}[:\-]?$")


def _is_header_text(v: Cell) -> bool:
    if not isinstance(v, str):
        return False
    if len(v) > 60 or re.fullmatch(r"[-+]?[\d.,/ ]+", v):
        return False
    return bool(re.search(r"[A-Za-z]", v))


def _header_score(rows: list[list[Cell]], idx: int, width: int) -> float:
    row = rows[idx]
    nonempty = [v for v in row if not is_empty(v)]
    if not nonempty:
        return -1
    texts = [v for v in nonempty if _is_header_text(v)]
    if len(texts) < max(1, math.ceil(0.6 * len(nonempty))):
        return -1
    # values below should be denser in numbers/dates than the header row and fill the same columns
    below = [r for r in rows[idx + 1: idx + 31] if any(not is_empty(v) for v in r)]
    if not below:
        return -1
    covered = sum(1 for c, v in enumerate(row) if not is_empty(v))
    fill_below = sum(sum(1 for c, v in enumerate(r) if not is_empty(v) and not is_empty(row[c])) for r in below)
    density = fill_below / (len(below) * max(covered, 1))
    typed_below = sum(1 for r in below for v in r if isinstance(v, (int, float, datetime))) / max(1, sum(
        1 for r in below for v in r if not is_empty(v)))
    distinct = len({str(t).lower() for t in texts}) / len(texts)
    return (len(texts) / width) * 2.0 + density + 0.6 * typed_below + 0.3 * distinct - idx * 0.01


def _dedupe(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for n in names:
        k = n.lower()
        if k in seen:
            seen[k] += 1
            out.append(f"{n} ({seen[k]})")
        else:
            seen[k] = 1
            out.append(n)
    return out


def segment_sheet(sheet: RawSheet) -> list[RawTable]:
    tables: list[RawTable] = []
    for b_idx, (c0, c1) in enumerate(_col_blocks(sheet.grid, sheet.merged)):
        width = c1 - c0 + 1
        sub = [(r_i, [row[c] if c < len(row) else None for c in range(c0, c1 + 1)]) for r_i, row in enumerate(sheet.grid)]
        sub = [(r_i, r) for r_i, r in sub if any(not is_empty(v) for v in r)]
        if not sub:
            continue
        rows = [r for _, r in sub]
        rnums = [r_i for r_i, _ in sub]

        # ---- header detection
        best, best_score = 0, -1.0
        for i in range(min(25, len(rows) - 1)):
            s = _header_score(rows, i, width)
            if s > best_score:
                best, best_score = i, s
        if width == 1:
            best = 0 if _is_header_text(rows[0][0]) and len(rows) > 1 else -1
        header_vals = rows[best] if best >= 0 else [None] * width
        letters = [get_column_letter(c + 1) for c in range(c0, c1 + 1)]
        names = [str(v) if not is_empty(v) else f"Column {letters[i]}" for i, v in enumerate(header_vals)]
        names = _dedupe(names)

        data, data_rnums = [], []
        threshold = 1 if width <= 2 else max(2, math.ceil(0.12 * width))
        for r, rn in zip(rows[best + 1:], rnums[best + 1:]):
            meaningful = sum(1 for v in r if not is_empty(v) and v != 0)
            if meaningful >= threshold:
                data.append(r)
                data_rnums.append(rn + 1)
        data, data_rnums = _drop_prefilled_tail(data, data_rnums, width)

        kind ="form" if _looks_like_form(rows, best, data, width, header_vals) else "table"
        r0, r1 = rnums[0], rnums[-1]
        t = RawTable(
            sheet=sheet.name, block_index=b_idx, kind=kind, first_col=c0, last_col=c1,
            header_row=(rnums[best] + 1) if best >= 0 else 0, columns=names, letters=letters,
            rows=data, row_numbers=data_rnums,
            range_ref=f"{letters[0]}{r0 + 1}:{letters[-1]}{r1 + 1}",
        )
        if kind == "form":
            t.form_cells = [(r_i - r0, c - c0, v) for r_i, r in sub for c, v in zip(range(c0, c1 + 1), r) if not is_empty(v)]
            t.merged = [(a - r0, b - c0, c - r0, d - c0) for (a, b, c, d) in sheet.merged
                        if b >= c0 and d <= c1 and a >= r0 and c <= r1]
        tables.append(t)
    return tables


def _sequence_columns(rows: list[list[Cell]], width: int) -> set[int]:
    """Columns that behave like auto-increment serials (value = previous + 1 almost everywhere)."""
    seq = set()
    for c in range(width):
        pairs = hits = 0
        prev = None
        for r in rows:
            v = r[c]
            if isinstance(v, int) and isinstance(prev, int):
                pairs += 1
                hits += (v - prev == 1)
            prev = v
        if pairs >= 10 and hits / pairs >= 0.9:
            seq.add(c)
    return seq


def _drop_prefilled_tail(data, rnums, width):
    """Spreadsheets are often pre-filled far below the real data with serial numbers and default
    values / formula results. Such rows carry no information beyond a serial, so the sparse region
    that trails the last properly-filled row is removed (serial columns are ignored when judging)."""
    if len(data) < 20:
        return data, rnums
    seq = _sequence_columns(data, width)
    sigs = [tuple(v for c, v in enumerate(r) if c not in seq) for r in data]
    fill = [sum(1 for v in s if not is_empty(v)) for s in sigs]
    typical = sorted(fill)[int(len(fill) * 0.9)] or 1   # p90: robust even when filler rows dominate
    last_full = max((i for i, f in enumerate(fill) if f >= 0.5 * typical), default=-1)
    tail = range(last_full + 1, len(data))
    if len(tail) >= 3:
        end = last_full + 1
    else:
        end = len(data)
        while end > last_full + 1 and fill[end - 1] <= 2:
            end -= 1
    keep = [i for i in range(end) if fill[i] > 0]
    return [data[i] for i in keep], [rnums[i] for i in keep]


def _looks_like_form(rows, header_idx, data, width, header_vals) -> bool:
    if width < 4:
        return False
    sig = tuple(str(v).lower() for v in header_vals if not is_empty(v))
    if header_idx >= 0 and len(sig) >= 3:
        repeats = sum(1 for r in rows if tuple(str(v).lower() for v in r if not is_empty(v))[: len(sig)] == sig)
        if repeats >= 2:
            return True        # multi-copy print template (original / duplicate / triplicate)
    if len(data) <= 12:
        cells = [v for r in data for v in r if not is_empty(v)]
        labels = [v for v in cells if isinstance(v, str) and _LABELISH.match(v) and (v.isupper() or v.endswith(("-", ":")))]
        if cells and len(labels) / len(cells) >= 0.3:
            return True
    filled = sum(1 for r in data for v in r if not is_empty(v))
    return bool(data) and width >= 8 and filled / (len(data) * width) < 0.12


# ----------------------------------------------------------------------------- column typing
def infer_dtype(values: list[Cell]) -> str:
    vals = [v for v in values if not is_empty(v)]
    if not vals:
        return "empty"
    n = len(vals)
    dates = sum(isinstance(v, datetime) for v in vals)
    nums = sum(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals)
    ints = sum(isinstance(v, int) and not isinstance(v, bool) for v in vals)
    if dates / n >= 0.8:
        return "date"
    if nums / n >= 0.85:
        return "integer" if ints == nums else "number"
    return "text"


def to_json_value(v: Cell) -> Any:
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if v.time() == time(0, 0) else v.isoformat(timespec="minutes")
    return v


def column_stats(values: list[Cell], dtype: str) -> dict:
    vals = [v for v in values if not is_empty(v)]
    st: dict[str, Any] = {"filled": len(vals), "empty": len(values) - len(vals)}
    if not vals:
        return st
    counts: dict[str, int] = {}
    for v in vals:
        k = str(to_json_value(v))
        counts[k] = counts.get(k, 0) + 1
    st["distinct"] = len(counts)
    st["top"] = sorted(counts.items(), key=lambda kv: -kv[1])[:6]
    if dtype in ("integer", "number"):
        nums = [v for v in vals if isinstance(v, (int, float))]
        st.update(min=min(nums), max=max(nums), sum=round(sum(nums), 2))
    if dtype == "date":
        ds = [v for v in vals if isinstance(v, datetime)]
        st.update(min=to_json_value(min(ds)), max=to_json_value(max(ds)))
    return st
