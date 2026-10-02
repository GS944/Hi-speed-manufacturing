"""Printable documents: Annexure (job-work details for an order) and Route Card.

Each document is built once as a structured payload (also used by the on-screen print preview) and
rendered to PDF (reportlab) or Excel (openpyxl) in the same layout.
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from ..db import SessionLocal, get_setting
from . import workspace

ANNEXURE_COLUMNS = [
    ("sl", "Sl.", 7), ("doc_no", "DC No.", 11), ("issue_date", "Invoice / DC Date", 21), ("vendor", "Vendor", 34),
    ("description", "Item Requested", 40), ("material", "Material", 15), ("qty", "Qty", 9),
    ("operation", "Operation", 30), ("unit_cost", "Unit Cost (Rs.)", 15), ("amount", "Amount (Rs.)", 17),
    ("value_of_goods", "Value of Goods (Rs.)", 18), ("hrc_spec", "HRC Spec", 15), ("hrc_checked", "HRC Checked", 17),
    ("hrc_result", "HRC Check", 19), ("inward_date", "Inward Date", 18),
]


def _fmt_date(d) -> str:
    if not d:
        return ""
    if isinstance(d, str):
        try:
            d = date.fromisoformat(d[:10])
        except ValueError:
            return d
    return d.strftime("%d-%m-%Y")


def _q(v) -> str:
    return "" if v is None else f"{v:g}"


def _money(v) -> str:
    return "" if v in (None, "") else f"{v:,.2f}"


def company() -> dict:
    with SessionLocal() as db:
        c = get_setting(db, "company") or {}
        t = get_setting(db, "annexure_template") or {}
    return {"name": c.get("name", "Your Company"), "address": c.get("address", []), "gstin": c.get("gstin", ""),
            "short_name": c.get("short_name", ""), "copies": t.get("copies") or ["ORIGINAL"],
            "tariff_code": t.get("tariff_code", ""), "dc_title": t.get("title", "")}


def annexure_payload(order_key: str, entry_ids: list[int] | None = None) -> dict | None:
    ws = workspace.get()
    lines = ws.find(order_key)
    if not lines:
        return None
    entries = []
    for ln in lines:
        for e in ln["job_work"]:
            if entry_ids and e["id"] not in entry_ids:
                continue
            entries.append((ln, e))
    seen, rows = set(), []
    for ln, e in entries:
        if e["id"] in seen:
            continue
        seen.add(e["id"])
        amount = (e["unit_cost"] or 0) * (e["qty"] or 0) if e["unit_cost"] is not None and e["qty"] is not None else None
        rows.append({
            "id": e["id"], "sl": len(rows) + 1, "doc_no": e["doc_no"], "issue_date": _fmt_date(e["issue_date"]),
            "vendor": e["vendor"] or "", "description": e["description"] or ln["roles"].get("item_description") or "",
            "material": e["material"] or ln["roles"].get("material") or "", "qty": e["qty"],
            "operation": e["operation"] or "", "unit_cost": e["unit_cost"], "amount": amount,
            "value_of_goods": e["value_of_goods"], "hrc_spec": e["hrc_spec"] or ln["roles"].get("hardness_spec") or "",
            "hrc_checked": e["hrc_checked"] or "",
            "hrc_result": {"ok": "OK", "out_of_spec": "OUT OF SPEC"}.get(e["hrc_result"], "Pending" if not e["hrc_checked"] else ""),
            "inward_date": _fmt_date(e["inward_date"]), "state": e["state"],
        })
    ln = lines[0]
    co = company()
    return {
        "type": "annexure", "title": "ANNEXURE — JOB WORK DETAILS",
        "company": co, "order_no": ln["key"], "generated_on": date.today().strftime("%d-%m-%Y"),
        "order": {
            "Order No.": ln["key"], "Customer": f"{ln['customer']}" + (f" ({ln['roles'].get('customer_code')})" if ln['roles'].get('customer_code') else ""),
            "Customer PO": ln["roles"].get("customer_po") or "", "Order Date": _fmt_date(ln["dates"].get("order_date")),
            "Due Date": _fmt_date(ln["dates"].get("due_date")),
            "Items Requested": "; ".join(str(l["roles"].get("item_description") or "") for l in lines),
            "Order Qty": " / ".join(f"{l['num'].get('order_qty'):g}" for l in lines if l["num"].get("order_qty") is not None),
            "Material / Hardness": f"{ln['roles'].get('material') or ''}  {ln['roles'].get('hardness_spec') or ''}".strip(),
            "Status": "COMPLETED" if all(l["status"] == "completed" for l in lines) else f"LIVE — {ln['stage']}",
        },
        "columns": [{"key": k, "label": l} for k, l, _ in ANNEXURE_COLUMNS],
        "rows": rows,
        "totals": {"qty": sum(r["qty"] or 0 for r in rows), "amount": round(sum(r["amount"] or 0 for r in rows), 2),
                   "value_of_goods": round(sum(r["value_of_goods"] or 0 for r in rows), 2), "entries": len(rows)},
    }


# ------------------------------------------------------------------------------------- PDF
def _styles():
    return {
        "h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=14, leading=17, textColor=colors.HexColor("#0f172a")),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11, leading=14, alignment=1),
        "sm": ParagraphStyle("sm", fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#334155")),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=7.2, leading=8.6),
        "cellb": ParagraphStyle("cellb", fontName="Helvetica-Bold", fontSize=7.2, leading=8.6),
        "copy": ParagraphStyle("copy", fontName="Helvetica-Bold", fontSize=8, alignment=2, textColor=colors.HexColor("#475569")),
    }


def _header(p: dict, st, copy_label: str, width: float):
    co = p["company"]
    left = [Paragraph(co["name"], st["h1"])] + [Paragraph(a, st["sm"]) for a in co["address"]]
    if co.get("gstin"):
        left.append(Paragraph(f"GSTIN: {co['gstin']}", st["sm"]))
    right = [Paragraph(copy_label, st["copy"]), Paragraph(f"Generated: {p['generated_on']}", st["copy"])]
    if p.get("doc_no"):
        right.append(Paragraph(f"Doc No: {p['doc_no']}", st["copy"]))
    t = Table([[left, right]], colWidths=[width * 0.7, width * 0.3])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, 0), 1.2, colors.HexColor("#0f172a")),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return t


def _meta_table(meta: dict, st, width: float):
    items = list(meta.items())
    rows = []
    for i in range(0, len(items), 3):
        row = []
        for k, v in items[i:i + 3]:
            row += [Paragraph(k, st["cellb"]), Paragraph(str(v), st["cell"])]
        row += [""] * (6 - len(row))
        rows.append(row)
    cw = width / 6
    t = Table(rows, colWidths=[cw * 0.75, cw * 1.25] * 3)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                           ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f5f9")),
                           ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#f1f5f9")),
                           ("BACKGROUND", (4, 0), (4, -1), colors.HexColor("#f1f5f9")),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    return t


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#64748b"))
    canvas.drawString(doc.leftMargin, 8 * mm, f"{getattr(doc, 'ot_title', '')}  ·  Order {getattr(doc, 'ot_order', '')}")
    canvas.drawRightString(doc.pagesize[0] - doc.rightMargin, 8 * mm, f"Page {doc.page}")
    canvas.restoreState()


def annexure_pdf(p: dict, copies: list[str] | None = None) -> bytes:
    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=14 * mm,
                            title=f"Annexure - Order {p['order_no']}", author=p["company"]["name"])
    doc.ot_title, doc.ot_order = p["title"], p["order_no"]
    st = _styles()
    width = page[0] - 20 * mm
    total_w = sum(w for *_, w in ANNEXURE_COLUMNS)
    col_w = [width * w / total_w for *_, w in ANNEXURE_COLUMNS]
    story = []
    for ci, copy_label in enumerate(copies or ["ORIGINAL"]):
        if ci:
            story.append(PageBreak())
        story += [_header(p, st, copy_label, width), Spacer(1, 5), Paragraph(p["title"], st["h2"]), Spacer(1, 5),
                  _meta_table(p["order"], st, width), Spacer(1, 7)]
        head = [Paragraph(l, st["cellb"]) for _, l, _ in ANNEXURE_COLUMNS]
        body = []
        for r in p["rows"]:
            row = []
            for k, _, _ in ANNEXURE_COLUMNS:
                v = r.get(k)
                v = _money(v) if k in ("unit_cost", "amount", "value_of_goods") else ("" if v is None else (f"{v:g}" if isinstance(v, float) else str(v)))
                row.append(Paragraph(v, st["cell"]))
            body.append(row)
        if not body:
            body = [[Paragraph("No job-work (annexure) entries recorded for this order yet.", st["cell"])] + [""] * (len(ANNEXURE_COLUMNS) - 1)]
        tot = [""] * len(ANNEXURE_COLUMNS)
        idx = {k: i for i, (k, _, _) in enumerate(ANNEXURE_COLUMNS)}
        tot[idx["description"]] = Paragraph("TOTAL", st["cellb"])
        tot[idx["qty"]] = Paragraph(f"{p['totals']['qty']:g}", st["cellb"])
        tot[idx["amount"]] = Paragraph(_money(p["totals"]["amount"]), st["cellb"])
        tot[idx["value_of_goods"]] = Paragraph(_money(p["totals"]["value_of_goods"]), st["cellb"])
        t = Table([head] + body + [tot], colWidths=col_w, repeatRows=1)
        style = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94a3b8")),
                 ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                 ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f1f5f9")),
                 ("VALIGN", (0, 0), (-1, -1), "TOP")]
        if not p["rows"]:
            style.append(("SPAN", (0, 1), (-1, 1)))
        for i, r in enumerate(p["rows"], start=1):
            if r["hrc_result"] == "OUT OF SPEC":
                style.append(("BACKGROUND", (idx["hrc_result"], i), (idx["hrc_result"], i), colors.HexColor("#fee2e2")))
        t.setStyle(TableStyle(style))
        story.append(t)
        story.append(Spacer(1, 18))
        sig = Table([[Paragraph("Prepared by", st["sm"]), Paragraph("Checked by (QC / HRC)", st["sm"]), Paragraph("Authorised Signatory", st["sm"])]],
                    colWidths=[width / 3] * 3, rowHeights=[22])
        sig.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.6, colors.HexColor("#475569")), ("VALIGN", (0, 0), (-1, -1), "BOTTOM")]))
        story.append(KeepTogether([Spacer(1, 14), sig]))
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


def annexure_xlsx(p: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Annexure {p['order_no']}"[:31]
    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    n = len(ANNEXURE_COLUMNS)
    last = get_column_letter(n)
    co = p["company"]
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = co["name"]
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = ", ".join(co["address"]) + (f"   GSTIN: {co['gstin']}" if co.get("gstin") else "")
    ws["A2"].font = Font(size=9, color="475569")
    ws.merge_cells(f"A3:{last}3")
    ws["A3"] = p["title"]
    ws["A3"].font = Font(bold=True, size=12)
    ws["A3"].alignment = Alignment(horizontal="center")
    r = 5
    items = list(p["order"].items())
    for i in range(0, len(items), 3):
        c = 1
        for k, v in items[i:i + 3]:
            ws.cell(r, c, k).font = Font(bold=True, size=9)
            ws.cell(r, c).fill = PatternFill("solid", fgColor="F1F5F9")
            ws.merge_cells(start_row=r, start_column=c + 1, end_row=r, end_column=c + 4)
            ws.cell(r, c + 1, str(v)).font = Font(size=9)
            c += 5
        r += 1
    r += 1
    for j, (_, label, w) in enumerate(ANNEXURE_COLUMNS, start=1):
        cell = ws.cell(r, j, label)
        cell.font = Font(bold=True, size=9)
        cell.fill = PatternFill("solid", fgColor="E2E8F0")
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = max(8, w * 0.42)
    header_row = r
    for row in p["rows"]:
        r += 1
        for j, (k, _, _) in enumerate(ANNEXURE_COLUMNS, start=1):
            cell = ws.cell(r, j, row.get(k))
            cell.border = border
            cell.font = Font(size=9, color="B91C1C" if k == "hrc_result" and row.get(k) == "OUT OF SPEC" else "0F172A")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if k in ("unit_cost", "amount", "value_of_goods"):
                cell.number_format = "#,##0.00"
    r += 1
    idx = {k: i + 1 for i, (k, _, _) in enumerate(ANNEXURE_COLUMNS)}
    ws.cell(r, idx["description"], "TOTAL").font = Font(bold=True)
    for k in ("qty", "amount", "value_of_goods"):
        col = get_column_letter(idx[k])
        c = ws.cell(r, idx[k], f"=SUM({col}{header_row + 1}:{col}{r - 1})" if p["rows"] else 0)
        c.font = Font(bold=True)
        c.number_format = "#,##0.00" if k != "qty" else "General"
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ------------------------------------------------------------------------------------- Route card
def route_card_payload(order_key: str) -> dict | None:
    ws = workspace.get()
    lines = ws.find(order_key)
    if not lines:
        return None
    co = company()
    cards = []
    for ln in lines:
        r, n = ln["roles"], ln["num"]
        cards.append({
            "order_no": ln["key"], "type": r.get("order_type"), "request": r.get("priority"),
            "po_date": _fmt_date(ln["dates"].get("order_date")), "delivery_date": _fmt_date(ln["dates"].get("due_date")),
            "material": r.get("material"), "hardness": r.get("hardness_spec"), "customer_code": r.get("customer_code"),
            "customer": ln["customer"], "customer_po": r.get("customer_po"), "spares": r.get("category"),
            "item": r.get("item_description"), "order_qty": n.get("order_qty"), "issued_qty": n.get("issued_qty"),
            "reject_qty": n.get("reject_qty"), "despatch_qty": n.get("despatch_qty"),
            "operations": [{"operation": e["operation"], "vendor": e["vendor"], "dc": e["doc_no"],
                            "date": _fmt_date(e["issue_date"]), "inward": _fmt_date(e["inward_date"])} for e in ln["job_work"]],
        })
    return {"type": "route_card", "title": f"ROUTE CARD — {co['short_name'] or co['name']}", "company": co,
            "order_no": lines[0]["key"], "cards": cards, "generated_on": date.today().strftime("%d-%m-%Y")}


def route_card_pdf(p: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
                            title=f"Route card - Order {p['order_no']}")
    doc.ot_title, doc.ot_order = p["title"], p["order_no"]
    st = _styles()
    width = A4[0] - 24 * mm
    story = []
    for i, c in enumerate(p["cards"]):
        if i:
            story.append(PageBreak())
        story += [Paragraph(p["title"], st["h1"]), Spacer(1, 6)]
        meta = {"Number": c["order_no"], "Type": c["type"] or "", "Request": c["request"] or "",
                "PO Date": c["po_date"], "Delivery Date": c["delivery_date"], "Customer Code": c["customer_code"] or "",
                "Customer": c["customer"], "Customer PO": c["customer_po"] or "", "Spares": c["spares"] or "",
                "Material": c["material"] or "", "Hardness": c["hardness"] or "", "Item": c["item"] or ""}
        story += [_meta_table(meta, st, width), Spacer(1, 8)]
        q = Table([[Paragraph(x, st["cellb"]) for x in ("Order Qty", "Issued", "Prod Rej", "Insp Rej", "Actual / Despatch")],
                   [_q(c["order_qty"]), _q(c["issued_qty"]), _q(c["reject_qty"]), "", _q(c["despatch_qty"])]],
                  colWidths=[width / 5] * 5, rowHeights=[16, 22])
        q.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94a3b8")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0"))]))
        story += [q, Spacer(1, 8)]
        stages = ["Cutting", "Turn / BM / AM", "HT", "CG / SG", "PKT"] + [o["operation"] for o in c["operations"] if o["operation"]]
        seen, rows = set(), [[Paragraph(x, st["cellb"]) for x in ("Operation", "Vendor / Machine", "DC / Date", "Operator", "Supervisor", "Inspector", "Qty")]]
        for s in stages:
            if s.upper() in seen:
                continue
            seen.add(s.upper())
            op = next((o for o in c["operations"] if o["operation"] == s), None)
            rows.append([Paragraph(s, st["cell"]), Paragraph(op["vendor"] if op else "", st["cell"]),
                         Paragraph(f"{op['dc']} / {op['date']}" if op else "", st["cell"]), "", "", "", ""])
        t = Table(rows, colWidths=[width * x for x in (0.2, 0.22, 0.16, 0.11, 0.11, 0.11, 0.09)], rowHeights=[16] + [22] * (len(rows) - 1))
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94a3b8")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story += [t, Spacer(1, 10), Paragraph("Remarks:", st["cellb"]), Spacer(1, 40)]
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()
